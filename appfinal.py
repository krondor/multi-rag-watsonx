import os
from langchain.document_loaders import PyPDFLoader, TextLoader
from langchain.indexes import VectorstoreIndexCreator
from langchain.chains import RetrievalQA
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.embeddings import HuggingFaceEmbeddings
from langchain.chains import LLMChain
from langchain.prompts import PromptTemplate
import streamlit as st
import tempfile

from ibm_watson_machine_learning.foundation_models import Model
from ibm_watson_machine_learning.foundation_models.extensions.langchain import WatsonxLLM
from ibm_watson_machine_learning.metanames import GenTextParamsMetaNames as GenParams
from ibm_watson_machine_learning.foundation_models.utils.enums import DecodingMethods

from pptx import Presentation
from docx import Document

# Initialize index to None
index = None

# Custom loader for DOCX files
class DocxLoader:
    def __init__(self, file_path):
        self.file_path = file_path

    def load(self):
        document = Document(self.file_path)
        text_content = []
        for para in document.paragraphs:
            text_content.append(para.text)
        return " ".join(text_content)

# Custom loader for PPTX files
class PptxLoader:
    def __init__(self, file_path):
        self.file_path = file_path

    def load(self):
        presentation = Presentation(self.file_path)
        text_content = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_content.append(shape.text)
        return " ".join(text_content)

# Caching function to load various file types
@st.cache_resource
def load_file(file_name, file_type):
    loaders = []
    
    if file_type == "pdf":
        loaders = [PyPDFLoader(file_name)]
    elif file_type == "docx":
        loader = DocxLoader(file_name)
        text = loader.load()
        
        with tempfile.NamedTemporaryFile(delete=False, suffix=".txt") as temp_file:
            temp_file.write(text.encode("utf-8"))
            temp_file_path = temp_file.name
        loaders = [TextLoader(temp_file_path)]
        
    elif file_type == "txt":
        loaders = [TextLoader(file_name)]
        
    elif file_type == "pptx":
        loader = PptxLoader(file_name)
        text = loader.load()
        
        with tempfile.NamedTemporaryFile(delete=False, suffix=".txt") as temp_file:
            temp_file.write(text.encode("utf-8"))
            temp_file_path = temp_file.name
        loaders = [TextLoader(temp_file_path)]
        
    else:
        st.error("Unsupported file type.")
        return None

    index = VectorstoreIndexCreator(
        embedding=HuggingFaceEmbeddings(model_name="all-MiniLM-L12-v2"),
        text_splitter=RecursiveCharacterTextSplitter(chunk_size=450, chunk_overlap=50)
    ).from_loaders(loaders)
    return index

def format_history():
    return ""

# Watsonx API setup using environment variables
watsonx_api_key = os.getenv("WATSONX_API_KEY")
watsonx_project_id = os.getenv("WATSONX_PROJECT_ID")

if not watsonx_api_key or not watsonx_project_id:
    st.error("API Key or Project ID is not set. Please set them as environment variables.")

prompt_template_br = PromptTemplate(
    input_variables=["context", "question"], 
    template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
I am a helpful assistant.

<|eot_id|>
{context}
<|start_header_id|>user<|end_header_id|>
{question}<|eot_id|>
"""
)

with st.sidebar:
    st.title("Watsonx RAG with Multiple docs")
    watsonx_model = st.selectbox("Model", ["meta-llama/llama-3-405b-instruct", "codellama/codellama-34b-instruct-hf", "ibm/granite-20b-multilingual"])
    max_new_tokens = st.slider("Max output tokens", min_value=100, max_value=4000, value=600, step=100)
    decoding_method = st.radio("Decoding", (DecodingMethods.GREEDY.value, DecodingMethods.SAMPLE.value))
    parameters = {
        GenParams.DECODING_METHOD: decoding_method,
        GenParams.MAX_NEW_TOKENS: max_new_tokens,
        GenParams.MIN_NEW_TOKENS: 1,
        GenParams.TEMPERATURE: 0,
        GenParams.TOP_K: 50,
        GenParams.TOP_P: 1,
        GenParams.STOP_SEQUENCES: [],
        GenParams.REPETITION_PENALTY: 1
    }
    st.info("Upload a PDF, DOCX, TXT, or PPTX file to use RAG")
    uploaded_file = st.file_uploader("Upload file", accept_multiple_files=False, type=["pdf", "docx", "txt", "pptx"])
    if uploaded_file is not None:
        bytes_data = uploaded_file.read()
        st.write("Filename:", uploaded_file.name)
        
        with open(uploaded_file.name, 'wb') as f:
            f.write(bytes_data)

        file_type = uploaded_file.name.split('.')[-1].lower()
        index = load_file(uploaded_file.name, file_type)

    model_name = watsonx_model

    def clear_messages():
        st.session_state.messages = []

    st.button('Clear messages', on_click=clear_messages)

st.info("Setting up Watsonx...")

my_credentials = {
    "url": "https://us-south.ml.cloud.ibm.com",
    "apikey": watsonx_api_key
}
params = parameters
project_id = watsonx_project_id
space_id = None
verify = False
model = WatsonxLLM(model=Model(model_name, my_credentials, params, project_id, space_id, verify))

if model:
    st.info(f"Model {model_name} ready.")
    chain = LLMChain(llm=model, prompt=prompt_template_br, verbose=True)

if chain:
    st.info("Chat ready.")
    if index:
        rag_chain = RetrievalQA.from_chain_type(
            llm=model,
            chain_type="stuff",
            retriever=index.vectorstore.as_retriever(),
            chain_type_kwargs={"prompt": prompt_template_br},
            return_source_documents=False,
            verbose=True
        )
        st.info("Chat with document ready.")

if "messages" not in st.session_state:
    st.session_state.messages = []

for message in st.session_state.messages:
    st.chat_message(message["role"]).markdown(message["content"])

prompt = st.chat_input("Ask your question here", disabled=False if chain else True)

if prompt:
    st.chat_message("user").markdown(prompt)

    response_text = None
    if rag_chain:
        response_text = rag_chain.run(prompt).strip()
    
    if not response_text:
        response = chain.run(question=prompt, context=format_history())
        response_text = response.strip("<|start_header_id|>assistant<|end_header_id|>").strip("<|eot_id|>")
        
    st.session_state.messages.append({'role': 'User', 'content': prompt })
    st.chat_message("assistant").markdown(response_text)
    st.session_state.messages.append({'role': 'Assistant', 'content': response_text })
